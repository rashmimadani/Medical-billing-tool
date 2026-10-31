import re
import shutil
import subprocess
import zipfile
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill

from invoice_service import build_invoice_number as strict_build_invoice_number, clean_client_row, clear_unused_service_rows, get_next_invoice_number, render_service_table
from tracker_service import ensure_monthly_tax_review_tracker


SERVICE_RULES = {
    "ASSESSMENT": {"aliases": {"ASSESSMENT", "A"}, "rates": {30: 2500, 60: 3750}},
    "PM": {"aliases": {"TPR", "PM"}, "rates": {30: 3000, 60: 4750}},
    "PT": {"aliases": {"PT", "REHAB"}, "rates": {30: 2500, 60: 3750}},
}

SERVICE_LABELS = {
    "ASSESSMENT": "Assessment",
    "PM": "Physical Medicine",
    "PT": "Rehabilitation",
}

SERVICE_ROW_MAP = {"ASSESSMENT": 15, "PM": 16, "PT": 17}
ALIASES = {}
for service, info in SERVICE_RULES.items():
    for code in info["aliases"]:
        ALIASES[code.upper()] = service


def normalize_service_code(code: str) -> str:
    cleaned = re.sub(r"[^A-Za-z]", "", str(code)).upper()
    if not cleaned:
        raise ValueError("Missing service code in session entry.")
    if cleaned not in ALIASES:
        raise ValueError(f"Unsupported service code '{code}'. Allowed codes: {sorted(ALIASES.keys())}.")
    return ALIASES[cleaned]


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value).lower())


def normalize_patient_name(value):
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_month(value):
    if pd.isna(value):
        return ""
    return str(value).strip()


def resolve_exact_file(filename: str):
    search_roots = [
        Path.cwd(),
        Path.cwd() / "data",
        Path.home(),
        Path.home() / "Downloads",
        Path("/home/rashmimadani/medical_tool"),
        Path("/home/rashmimadani/medical_tool/data"),
        Path("/mnt/c/Users/rmadani/Downloads"),
        Path("/mnt/c/Users"),
        Path("/mnt"),
    ]

    for root in search_roots:
        candidate = root / filename
        if candidate.exists():
            return candidate

    for root in [Path.home(), Path("/mnt/c/Users/rmadani/Downloads"), Path("/mnt/c/Users")]:
        if not root.exists():
            continue
        try:
            matches = list(root.glob(filename)) + list(root.glob(f"*{filename}*"))
        except PermissionError:
            continue
        if matches:
            unique_matches = []
            seen = set()
            for match in matches:
                if match not in seen:
                    unique_matches.append(match)
                    seen.add(match)
            return unique_matches[0]

    return Path(filename)


def parse_discount_entries(raw_text: str):
    discounts = {}
    warnings = []
    if not raw_text or not str(raw_text).strip():
        return discounts, warnings

    for line_no, raw_line in enumerate(str(raw_text).splitlines(), start=1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue

        match = re.match(r"^(?P<patient>.+?)\s*(?:-|:)\s*(?P<percent>\d+(?:\.\d+)?)\s*%\s*(?:discount)?\s*$", line, flags=re.IGNORECASE)
        if not match:
            match = re.match(r"^(?P<patient>.+?)\s+(?P<percent>\d+(?:\.\d+)?)\s*%\s*(?:discount)?\s*$", line, flags=re.IGNORECASE)
        if not match:
            warnings.append(f"Line {line_no}: Could not parse discount entry '{line}'. Expected format: Name - 20% discount.")
            continue

        patient = normalize_patient_name(match.group("patient"))
        percent = float(match.group("percent"))
        if percent < 0 or percent > 100:
            warnings.append(f"Line {line_no}: Discount for '{patient}' must be between 0% and 100%.")
            continue
        discounts[patient.casefold()] = float(percent)

    return discounts, warnings


def apply_discount_to_patient_totals(patient_totals: pd.DataFrame, discount_map: dict[str, float] | None = None):
    working = patient_totals.copy()
    if working.empty:
        working["gross_amount"] = []
        working["discount_percent"] = []
        working["net_amount"] = []
        working["amount"] = []
        return working

    discount_map = discount_map or {}
    working["gross_amount"] = working["total_amount"].astype(int)
    working["discount_percent"] = working["patient"].map(lambda patient: float(discount_map.get(normalize_patient_name(patient).casefold(), 0.0)))
    working["net_amount"] = (working["gross_amount"] * (1 - working["discount_percent"] / 100)).round().astype(int)
    working["amount"] = working["net_amount"]
    return working


def parse_monthly_log(raw_text: str):
    rows = []
    warnings = []
    lines = raw_text.splitlines()

    for line_no, raw_line in enumerate(lines, start=1):
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if not re.search(r"[A-Za-z]", line):
            continue

        match = re.match(r"^(?P<patient>.+?)\s*-\s*(?P<items>.+)$", line)
        if not match:
            warnings.append(f"Line {line_no}: Could not parse line '{line}'. Expected format: Name - 1(30)(PT) + 2(60)(TPR).")
            continue

        patient = re.sub(r"\s+", " ", match.group("patient")).strip()
        items_expression = match.group("items").strip()
        if not items_expression:
            warnings.append(f"Line {line_no}: Missing session items for patient '{patient}'.")
            continue

        try:
            entry_items = [part.strip() for part in re.split(r"\s*\+\s*", items_expression) if part.strip()]
            if not entry_items:
                raise ValueError("No session entries found.")

            for entry in entry_items:
                match_item = re.match(
                    r"^(?P<count>\d+)\s*\(\s*(?P<duration>\d+)\s*\)\s*(?:\(\s*(?P<code>[^)]+)\s*\)|\s*(?P<code2>[A-Za-z]+))?\s*$",
                    entry,
                    flags=re.IGNORECASE,
                )
                if not match_item:
                    raise ValueError(f"Unrecognized session entry '{entry}'.")

                count = int(match_item.group("count"))
                duration = int(match_item.group("duration"))
                if duration == 45:
                    duration = 30
                service_code = normalize_service_code(match_item.group("code") or match_item.group("code2"))
                service_rate = SERVICE_RULES[service_code]["rates"][duration]
                amount = service_rate * count

                rows.append(
                    {
                        "patient": patient,
                        "session_count": count,
                        "duration_minutes": duration,
                        "service_code": service_code,
                        "service_name": SERVICE_LABELS[service_code],
                        "rate": service_rate,
                        "amount": amount,
                    }
                )
        except ValueError as exc:
            warnings.append(f"Line {line_no}: {exc}")

    parsed_df = pd.DataFrame(rows)
    if not parsed_df.empty:
        parsed_df = parsed_df.sort_values(["patient", "service_code"]).reset_index(drop=True)
    return parsed_df, warnings


def enrich_service_rows_with_discount(parsed_df: pd.DataFrame, discount_map: dict[str, float] | None = None):
    display_df = parsed_df.copy()
    if display_df.empty:
        return display_df

    discount_map = discount_map or {}
    display_df["gross_amount"] = display_df["amount"].astype(int)
    display_df["discount_percent"] = display_df["patient"].map(
        lambda patient: float(discount_map.get(normalize_patient_name(str(patient)).casefold(), 0.0))
    )
    display_df["net_amount"] = (display_df["gross_amount"] * (1 - display_df["discount_percent"] / 100)).round().astype(int)
    return display_df


def summarize_patient_totals(parsed_df: pd.DataFrame):
    if parsed_df.empty:
        return pd.DataFrame(columns=["patient", "service_summary", "total_sessions", "total_amount", "service_breakdown"])

    grouped = (
        parsed_df.groupby(["patient", "service_code", "service_name"], as_index=False)
        .agg(service_sessions=("session_count", "sum"), total_amount=("amount", "sum"), total_duration=("duration_minutes", "sum"))
        .sort_values(["patient", "service_code"])
        .reset_index(drop=True)
    )

    patient_totals = (
        parsed_df.groupby("patient", as_index=False)
        .agg(total_sessions=("session_count", "sum"), total_amount=("amount", "sum"))
        .sort_values("patient")
        .reset_index(drop=True)
    )

    patient_totals["service_summary"] = patient_totals["patient"].map(
        grouped.groupby("patient")["service_name"].agg(lambda names: ", ".join(sorted(set(names))))
    )
    patient_totals["service_breakdown"] = patient_totals["patient"].map(
        grouped.groupby("patient").apply(
            lambda sub: "; ".join(
                f"{row['service_name']}: {int(row['service_sessions'])} session(s) x {int(row['total_duration'])} min"
                for _, row in sub.iterrows()
            )
        )
    )
    return patient_totals


def build_invoice_number(patient_name: str, month_label: str, patient_index: int = 1, prefix: str = "2627", series_code: str = "NL", start_sequence: int = 16):
    return strict_build_invoice_number(patient_name, month_label, patient_index, prefix=prefix, series_code=series_code, start_sequence=start_sequence)


def set_invoice_header(ws, patient_name: str, invoice_number: str, invoice_date):
    if isinstance(invoice_date, datetime):
        cell_value = invoice_date.date()
    elif isinstance(invoice_date, date):
        cell_value = invoice_date
    elif isinstance(invoice_date, str):
        parsed = None
        for fmt in ("%d-%b-%Y", "%Y-%m-%d", "%d-%B-%Y", "%d/%m/%Y"):
            try:
                parsed = datetime.strptime(invoice_date, fmt).date()
                break
            except ValueError:
                continue
        cell_value = parsed if parsed is not None else invoice_date
    else:
        cell_value = invoice_date

    ws["I3"] = invoice_number
    ws["I4"] = cell_value
    if isinstance(cell_value, (date, datetime)):
        ws["I4"].number_format = "DD-MMM-YYYY"

    clean_client_row(ws, patient_name)


def clear_service_row(ws, row_number: int):
    for col in range(1, 9):
        cell = ws.cell(row=row_number, column=col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = None
        cell.border = Border()


def force_single_page_layout(ws):
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def apply_exact_invoice_template(workbook, patient_name: str, invoice_date: date, invoice_number: str, patient_df: pd.DataFrame, discount_percent: float = 0.0):
    ws = workbook.active
    force_single_page_layout(ws)
    set_invoice_header(ws, patient_name, invoice_number, invoice_date)
    ws["A6"] = None
    ws["A11"] = None

    if patient_df.empty:
        render_service_table(ws, [], discount_percent=discount_percent)
        return

    service_totals = (
        patient_df.groupby("service_code", as_index=False)
        .agg(total_sessions=("session_count", "sum"), duration_minutes=("duration_minutes", "max"), rate=("rate", "max"), total_amount=("amount", "sum"))
        .sort_values("service_code")
        .reset_index(drop=True)
    )

    ordered_services = ["ASSESSMENT", "PM", "PT"]
    active_service_rows = []
    for service_code in ordered_services:
        if service_code not in set(service_totals["service_code"]):
            continue
        service_entry = service_totals[service_totals["service_code"] == service_code].iloc[0]
        quantity = int(service_entry["total_sessions"])
        if quantity <= 0:
            continue
        duration = int(service_entry["duration_minutes"])
        rate = int(service_entry["rate"])
        label = SERVICE_LABELS[service_code]
        active_service_rows.append({
            "description": f"{label} ({duration} min)",
            "rate": rate,
            "quantity": quantity,
        })

    render_service_table(ws, active_service_rows, discount_percent=discount_percent)


def lookup_row_for_service(service_code: str):
    return SERVICE_ROW_MAP[service_code]


def convert_xlsx_to_pdf(xlsx_path: Path, pdf_path: Path):
    soffice_candidates = [
        shutil.which("soffice"),
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        "/usr/lib/libreoffice/program/soffice",
        "/snap/bin/libreoffice",
    ]
    soffice_path = next((candidate for candidate in soffice_candidates if candidate and Path(candidate).exists()), None)
    if not soffice_path:
        raise RuntimeError("LibreOffice CLI 'soffice' is required but was not found in PATH or the standard install locations.")

    command = [
        soffice_path,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(pdf_path.parent),
        str(xlsx_path),
    ]
    completed = subprocess.run(command, capture_output=True, text=True)
    if completed.returncode != 0:
        raise RuntimeError(f"PDF conversion failed for {xlsx_path}: {completed.stderr.strip() or completed.stdout.strip()}")

    expected_pdf = xlsx_path.with_suffix(".pdf")
    if expected_pdf.exists() and expected_pdf != pdf_path:
        expected_pdf.replace(pdf_path)


def update_income_tracker(tracker_path: Path, patient_name: str, month_label: str, amount: float, total_sessions: int, service_summary: str, invoice_number: str):
    tracker_path = Path(tracker_path)
    if tracker_path.exists():
        wb = load_workbook(tracker_path)
        ws = wb.active
        force_single_page_layout(ws)
        next_row = ws.max_row + 1
        entry_id = ws.max_row if ws["B2"].value is not None else 1
        ws[f"B{next_row}"] = int(entry_id)
        ws[f"C{next_row}"] = patient_name
        ws[f"D{next_row}"] = "Consultation"
        ws[f"E{next_row}"] = "One-time"
        ws[f"F{next_row}"] = "N.A."
        ws[f"L{next_row}"] = amount
        ws[f"M{next_row}"] = "Bank"
        ws[f"O{next_row}"] = "Completed"
        ws[f"Q{next_row}"] = invoice_number
        ws[f"R{next_row}"] = "Sent"
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Payments"
        force_single_page_layout(ws)
        ws.append([None, "Entry ID", "Client Name", "Program/Consultation", "Billing Type", "Prepaid / Monthly / N.A. ", "Months Committed", "Months Paid", "Months remaining Payment", "Payment date", "Amount per month", "Amount (INR)", "Payment Method", "Next Due Date", None, "Status", "Invoice raised?", "Invoice #", "Status Details"])
        ws.append([None, 1, patient_name, "Consultation", "One-time", "N.A.", None, None, None, None, None, amount, "Bank", None, None, "Completed", None, invoice_number, "Sent"])
    wb.save(tracker_path)


def build_monthly_summary_workbook(month_label: str, invoice_rows, summary_path: Path):
    summary_path = Path(summary_path)
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Invoice Summary"
    summary_ws.freeze_panes = "A3"
    summary_ws.sheet_view.showGridLines = False

    summary_ws.merge_cells("A1:I1")
    summary_ws["A1"] = f"Monthly Billing Summary - {month_label}"
    summary_ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary_ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    summary_ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["Month", "Patient", "Invoice Number", "Invoice Date", "Gross Amount", "Net Amount", "Total Sessions", "Service Breakdown", "Review Status"]
    summary_ws.append(headers)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in summary_ws[2]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_gross = 0
    total_net = 0
    total_sessions = 0
    for row in invoice_rows:
        gross = int(row.get("gross_amount", row.get("amount", 0)))
        net = int(row.get("net_amount", gross))
        sessions = int(row.get("total_sessions", 0))
        total_gross += gross
        total_net += net
        total_sessions += sessions

        service_detail = row.get("service_detail") or row.get("service_summary", "")
        if not service_detail and row.get("service_summary"):
            service_detail = row["service_summary"]

        summary_ws.append([
            month_label,
            row["patient"],
            row["invoice_number"],
            row.get("invoice_date", ""),
            gross,
            net,
            sessions,
            service_detail,
            row.get("review_status", "Reviewed"),
        ])

    for row in summary_ws.iter_rows(min_row=3, max_row=summary_ws.max_row, min_col=1, max_col=9):
        for cell in row:
            if cell.column in (5, 6):
                cell.number_format = '#,##0'
            if cell.column == 8:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column_cells in summary_ws.columns:
        cells = [cell for cell in column_cells if hasattr(cell, "column_letter")]
        if not cells:
            continue
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in cells)
        summary_ws.column_dimensions[cells[0].column_letter].width = min(max(length + 2, 12), 30)

    summary_ws.append([])
    total_row = ["", "", "", "Month Total", total_gross, total_net, total_sessions, "All services included above", ""]
    summary_ws.append(total_row)
    total_fill = PatternFill("solid", fgColor="2E7D32")
    for cell in summary_ws[summary_ws.max_row][3:7]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = total_fill
    summary_ws[summary_ws.max_row][7].font = Font(bold=True)
    summary_ws[summary_ws.max_row][7].fill = PatternFill("solid", fgColor="E2F0D9")
    summary_ws[summary_ws.max_row][7].alignment = Alignment(wrap_text=True)

    log_ws = wb.create_sheet("Invoice Log")
    log_ws.sheet_view.showGridLines = False
    log_ws.append(["Month", "Patient", "Invoice Number", "Gross Amount", "Net Amount", "Total Sessions", "Service Breakdown", "Review Status"])
    for cell in log_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")

    for row in invoice_rows:
        gross = int(row.get("gross_amount", row.get("amount", 0)))
        net = int(row.get("net_amount", gross))
        sessions = int(row.get("total_sessions", 0))
        log_ws.append([
            month_label,
            row["patient"],
            row["invoice_number"],
            gross,
            net,
            sessions,
            row.get("service_detail") or row.get("service_summary", ""),
            row.get("review_status", "Reviewed"),
        ])

    for col in ["D", "E"]:
        for cell in log_ws[col]:
            cell.number_format = '#,##0'
    for cell in log_ws["G"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for column_cells in log_ws.columns:
        cells = [cell for cell in column_cells if hasattr(cell, "column_letter")]
        if not cells:
            continue
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in cells)
        log_ws.column_dimensions[cells[0].column_letter].width = min(max(length + 2, 12), 30)

    summary_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(summary_path)
    return summary_path


def build_invoice_bundle_zip(invoice_files, summary_path: Path, bundle_path: Path):
    bundle_path = Path(bundle_path)
    with zipfile.ZipFile(bundle_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file_path in invoice_files:
            archive.write(file_path, arcname=file_path.name)
        archive.write(summary_path, arcname=summary_path.name)
    return bundle_path


def validate_monthly_data(parsed_df: pd.DataFrame, warnings: list[str]) -> list[str]:
    issues = list(warnings)
    if parsed_df.empty:
        issues.append("No valid session rows were parsed; nothing will be generated.")
        return issues

    if parsed_df["amount"].sum() <= 0:
        issues.append("The parsed session total is zero or negative; please review the input log.")

    empty_service_rows = parsed_df[parsed_df["service_name"].isna() | parsed_df["service_name"].astype(str).str.strip().eq("")]
    if not empty_service_rows.empty:
        issues.append("Some session rows are missing a valid service type.")

    duplicates = parsed_df.duplicated(subset=["patient", "service_code", "duration_minutes"]).any()
    if duplicates:
        issues.append("Duplicate patient/service entries were detected; review before generating invoices.")

    return issues


def generate_exact_invoice_bundle(parsed_df: pd.DataFrame, month_label: str, template_path: Path, output_dir: Path, tracker_path: Path, discount_text: str | dict | None = ""):
    if parsed_df.empty:
        raise ValueError("No parsed invoice rows available.")

    template_path = Path(template_path)
    if not template_path.exists():
        raise FileNotFoundError(
            "The required exact template file was not found. Place the workbook named "
            "'Physical_Medicine_Template invoice_Price Included_1-Jun-26.xlsx' in the workspace or the same directory."
        )

    output_dir.mkdir(parents=True, exist_ok=True)
    patient_totals = summarize_patient_totals(parsed_df)
    if isinstance(discount_text, dict):
        discount_map = {str(key).casefold(): float(value) for key, value in discount_text.items()}
        discount_warnings = []
    else:
        discount_map, discount_warnings = parse_discount_entries(discount_text)
    if discount_warnings:
        st.warning("\n".join(discount_warnings))
    patient_totals = apply_discount_to_patient_totals(patient_totals, discount_map)
    invoices = []
    tracker_rows = []

    for idx, patient in enumerate(patient_totals["patient"], start=1):
        patient_name = str(patient)
        patient_df = parsed_df[parsed_df["patient"] == patient_name].copy()
        invoice_number = build_invoice_number(patient_name, month_label, idx, prefix="2627", series_code="NL", start_sequence=16)
        invoice_date = datetime.strptime(f"{month_label} 01", "%b %Y %d").date()
        invoice_date = invoice_date.replace(day=1)
        last_day = pd.Timestamp(invoice_date).to_period("M").end_time.date()
        discount_pct = float(patient_totals.loc[patient_totals["patient"] == patient_name, "discount_percent"].iloc[0])
        gross_total = int(patient_totals.loc[patient_totals["patient"] == patient_name, "gross_amount"].iloc[0])
        net_total = int(patient_totals.loc[patient_totals["patient"] == patient_name, "net_amount"].iloc[0])

        source_template = load_workbook(template_path)
        apply_exact_invoice_template(source_template, patient_name, last_day, invoice_number, patient_df, discount_percent=discount_pct)

        excel_export_dir = Path("Invoices_Excel")
        excel_export_dir.mkdir(parents=True, exist_ok=True)
        output_xlsx = excel_export_dir / f"Invoice_{invoice_number}_{patient_name}.xlsx"
        source_template.save(output_xlsx)

        pdf_export_dir = Path("Invoices_PDF")
        pdf_export_dir.mkdir(parents=True, exist_ok=True)
        pdf_output = pdf_export_dir / f"Invoice_{invoice_number}_{patient_name}.pdf"
        convert_xlsx_to_pdf(output_xlsx, pdf_output)

        update_income_tracker(
            tracker_path,
            patient_name,
            month_label,
            net_total,
            int(patient_totals.loc[patient_totals["patient"] == patient_name, "total_sessions"].iloc[0]),
            patient_totals.loc[patient_totals["patient"] == patient_name, "service_summary"].iloc[0],
            invoice_number,
        )

        total_sessions = int(patient_totals.loc[patient_totals["patient"] == patient_name, "total_sessions"].iloc[0])
        service_summary = patient_totals.loc[patient_totals["patient"] == patient_name, "service_summary"].iloc[0]
        service_breakdown = patient_totals.loc[patient_totals["patient"] == patient_name, "service_breakdown"].iloc[0]
        tracker_rows.append(
            {
                "patient": patient_name,
                "invoice_number": invoice_number,
                "gross_amount": gross_total,
                "net_amount": net_total,
                "amount": net_total,
                "total_sessions": total_sessions,
                "service_summary": service_summary,
                "service_detail": service_breakdown,
                "review_status": "Reviewed",
                "invoice_date": last_day,
            }
        )
        invoices.append({"patient": patient_name, "invoice_number": invoice_number, "xlsx": output_xlsx, "pdf": pdf_output, "gross_amount": gross_total, "net_amount": net_total, "amount": net_total, "total_sessions": total_sessions, "service_summary": service_summary, "service_detail": service_breakdown, "invoice_date": last_day})

    ensure_monthly_tax_review_tracker(tracker_path, month_label, tracker_rows)

    summary_filename = f"Monthly_Summary_{month_label.replace(' ', '_')}.xlsx"
    summary_path = output_dir / summary_filename
    build_monthly_summary_workbook(month_label, tracker_rows, summary_path)

    bundle_filename = f"Invoice_Bundle_{month_label.replace(' ', '_')}.zip"
    bundle_path = output_dir / bundle_filename
    build_invoice_bundle_zip([invoice["xlsx"] for invoice in invoices] + [invoice["pdf"] for invoice in invoices], summary_path, bundle_path)

    st.session_state["monthly_summary_path"] = summary_path
    st.session_state["monthly_bundle_path"] = bundle_path
    return invoices


if __name__ == "__main__":
    st.set_page_config(page_title="Cell-Optimal Lyfe Billing Tool", layout="wide")
    st.title("Cell-Optimal Lyfe Medical Solutions")
    st.subheader("Monthly Billing & Invoicing Automation")

    with st.sidebar:
        st.markdown("### Configuration")
        month = st.text_input("Month (e.g. Jan 2026)", value="Jan 2026")
        payment_method = st.text_input("Default payment method", value="Bank")
        tracker_path = resolve_exact_file("CellOptimal_Income_Tracker_16Oct25.xlsx")
        template_path = resolve_exact_file("Physical_Medicine_Template invoice_Price Included_1-Jun-26.xlsx")
        output_dir = Path(st.text_input("Invoice output folder", value="data/invoices"))
        st.caption("This workflow locks to the exact invoice layout and cell mapping required for the template workbook.")

    st.write("Paste the raw monthly staff log below. The parser will normalize the entries and show the final billable amounts before saving anything.")
    raw_log = st.text_area(
        "Monthly staff log",
        value=(
            "Ravi - 7 (30)(PT)\n"
            "Abhinay - 1 (60) (PT)+ 7(30)(TPR)\n"
            "Punit - 7 (30)(TPR)\n"
            "Shilpa - 7 (30)(TPR)\n"
            "Gautham - 2(60)(TPR)\n"
            "Rashmi - 2(45)(PT)\n"
            "Kush-2(45)(TPR)\n"
            "George -4(30)(TPR)\n"
            "Simran - 6 (30) (TPR)\n"
            "Shanta - 2 (30)(PT)\n"
            "Dr kamala -7(30)(TPR)\n"
            "Bhanu- 5 (30)(TPR)\n"
            "Col Sushil - 1(30)(PT)\n"
            "Anu-1 (30)(PT)\n"
            "Naina- 4(30)(TPR)\n"
            "Shivaram-1(30)(TPR)\n"
            "Dr.arun kumar- 2(30)(TPR)\n"
        ),
        height=250,
    )
    discount_log = st.text_area(
        "Discounts for eligible patients (optional)",
        value="",
        height=120,
        placeholder="Abhinay - 20% discount\nShilpa - 25% discount",
    )

    if st.button("Parse log and review calculations"):
        parsed_df, warnings = parse_monthly_log(raw_log)
        discount_map, discount_warnings = parse_discount_entries(discount_log)
        if parsed_df.empty:
            st.error("No valid rows were parsed. Please check the staff log format and try again.")
            if warnings:
                st.warning("\n".join(warnings))
        else:
            patient_totals = summarize_patient_totals(parsed_df)
            patient_totals = apply_discount_to_patient_totals(patient_totals, discount_map)
            validation_issues = validate_monthly_data(parsed_df, warnings)
            if discount_warnings:
                validation_issues.extend(discount_warnings)
            st.success(f"Successfully parsed {len(parsed_df)} session entries across {len(patient_totals)} patients.")
            if validation_issues:
                st.warning("Validation checks:\n" + "\n".join(validation_issues))
            service_review_df = enrich_service_rows_with_discount(parsed_df, discount_map)
            st.subheader("Patient totals")
            st.dataframe(patient_totals[["patient", "total_sessions", "total_amount", "discount_percent", "net_amount"]])
            st.subheader("Service-by-service details")
            st.dataframe(service_review_df[["patient", "session_count", "duration_minutes", "service_code", "service_name", "rate", "amount", "gross_amount", "net_amount"]])
            st.session_state["parsed_df"] = parsed_df
            st.session_state["service_review_df"] = service_review_df
            st.session_state["patient_totals"] = patient_totals
            st.session_state["discount_map"] = discount_map
            st.session_state["warnings"] = validation_issues

    if "parsed_df" in st.session_state:
        parsed_df = st.session_state["parsed_df"]
        service_review_df = st.session_state.get("service_review_df", enrich_service_rows_with_discount(parsed_df, st.session_state.get("discount_map", {})))
        patient_totals = st.session_state["patient_totals"]
        warnings = st.session_state.get("warnings", [])

        st.subheader("Review & commit")
        st.write("Review the parsed totals below, then generate the invoice pack once you are happy with the month summary.")

        with st.container(border=True):
            for _, patient_row in patient_totals.iterrows():
                patient_name = str(patient_row["patient"])
                patient_number = build_invoice_number(patient_name, normalize_month(month), 1 if patient_name == patient_totals.iloc[0]["patient"] else list(patient_totals["patient"]).index(patient_name) + 1)
                net_amount = int(patient_row.get("net_amount", patient_row.get("total_amount", 0)))
                discount_percent = float(patient_row.get("discount_percent", 0.0))
                with st.expander(f"{patient_name} • {int(patient_row['total_sessions'])} sessions • ₹{net_amount:,} net", expanded=False):
                    st.write(patient_row["service_breakdown"])
                    if discount_percent > 0:
                        st.caption(f"Discount applied: {discount_percent:.0f}%")
                    st.caption(f"Estimated invoice reference: {patient_number}")

        with st.form("generate_bulk_invoices"):
            st.write("Ready to generate? This will build the exact invoice Excel files, convert them to PDF, append the tracker rows, and create the month summary export.")
            submitted = st.form_submit_button("Generate exact-template invoices", type="primary")

        if submitted:
            month_value = normalize_month(month)
            if not month_value:
                st.error("Month cannot be blank.")
            else:
                try:
                    invoices = generate_exact_invoice_bundle(parsed_df, month_value, template_path, output_dir, tracker_path, st.session_state.get("discount_map", {}))
                    st.success("Exact-template invoice workflow completed successfully.")

                    summary_path = st.session_state.get("monthly_summary_path")
                    bundle_path = st.session_state.get("monthly_bundle_path")

                    if summary_path and summary_path.exists():
                        with open(summary_path, "rb") as fh:
                            st.download_button(
                                label=f"Download month summary (.xlsx)",
                                data=fh.read(),
                                file_name=summary_path.name,
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            )

                    if bundle_path and bundle_path.exists():
                        with open(bundle_path, "rb") as fh:
                            st.download_button(
                                label=f"Download all Excel + PDF invoices (.zip)",
                                data=fh.read(),
                                file_name=bundle_path.name,
                                mime="application/zip",
                            )

                    for invoice in invoices:
                        st.markdown(f"### {invoice['patient']}")
                        with open(invoice["xlsx"], "rb") as fh:
                            st.download_button(
                                label=f"Download {invoice['patient']} invoice (.xlsx)",
                                data=fh.read(),
                                file_name=invoice["xlsx"].name,
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            )
                        with open(invoice["pdf"], "rb") as fh:
                            st.download_button(
                                label=f"Download {invoice['patient']} invoice (.pdf)",
                                data=fh.read(),
                                file_name=invoice["pdf"].name,
                                mime="application/pdf",
                            )
                except Exception as exc:  # pragma: no cover - UI error handling
                    st.error(f"The exact-template workflow could not complete. Details: {exc}")

        if warnings:
            st.warning("\n".join(warnings))
