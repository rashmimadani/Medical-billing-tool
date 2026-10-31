import re

from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side

SERIES_PREFIX = "2627"
SERIES_CODE = "NL"
SERIES_START = 16


def _unmerge_row_ranges(ws, start_row: int, end_row: int):
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row <= end_row and merged_range.max_row >= start_row:
            ws.unmerge_cells(str(merged_range))


def _safe_clear_cells(ws, row_start: int, row_end: int, col_start: int = 1, col_end: int = 8):
    for row in range(row_start, row_end + 1):
        for col in range(col_start, col_end + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None
            cell.border = Border()
            cell.alignment = Alignment(horizontal="left", vertical="center")


def _apply_full_table_border(ws, row: int, start_col: int = 1, end_col: int = 8):
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        if isinstance(cell, MergedCell):
            continue
        cell.border = border
        if col in {6, 7, 8}:
            cell.alignment = Alignment(horizontal="right", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")


def _apply_summary_currency_style(ws, row: int, value, *, percentage: bool = False, bold: bool = False):
    cell = ws[f"H{row}"]
    cell.value = value
    cell.number_format = "0%" if percentage else "#,##0"
    cell.font = Font(name="Calibri", size=11, bold=bold, italic=False)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    thin = Side(style="thin", color="000000")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws[f"A{row}"].font = Font(name="Calibri", size=11, bold=bold, italic=False)
    ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="center")
    return cell


def clean_client_row(ws, patient_name: str):
    """Keep a single clean, fully closed Client Name block in the template."""
    _unmerge_row_ranges(ws, 10, 12)
    _safe_clear_cells(ws, 10, 12, 1, 5)

    ws.merge_cells("A10:B10")
    ws.merge_cells("C10:E10")
    ws["A10"] = "Client Name"
    ws["C10"] = patient_name

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell_ref in ["A10", "C10", "D10", "E10"]:
        ws[cell_ref].border = border
        ws[cell_ref].alignment = Alignment(horizontal="left", vertical="center")
    ws["E10"].border = border

    for row in (11, 12):
        for col in range(1, 6):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell, MergedCell):
                continue
            cell.value = None
            cell.border = Border()
            cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[11].height = 0
    ws.row_dimensions[12].height = 0
    return ws


def render_service_table(ws, service_rows, discount_percent: float = 0.0):
    """Render only the active service rows and compact the summary immediately beneath them."""
    if service_rows is None:
        service_rows = []

    _unmerge_row_ranges(ws, 14, 22)
    _safe_clear_cells(ws, 14, 22, 1, 8)

    ws.merge_cells("A14:E14")
    ws["A14"] = "Description"
    ws["F14"] = "Amount / unit (₹)"
    ws["G14"] = "Unit quantity"
    ws["H14"] = "Total amount (₹)"
    for cell_ref in ["A14", "F14", "G14", "H14"]:
        ws[cell_ref].font = Font(name="Calibri", size=11, bold=True, italic=False)
        ws[cell_ref].border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
        ws[cell_ref].alignment = Alignment(horizontal="right" if cell_ref in {"F14", "G14", "H14"} else "left", vertical="center")

    active_rows = list(service_rows)
    for idx, service in enumerate(active_rows):
        row_num = 15 + idx
        ws.merge_cells(f"A{row_num}:E{row_num}")
        ws[f"A{row_num}"] = service["description"]
        ws[f"A{row_num}"].font = Font(name="Calibri", size=11, bold=False, italic=False)
        ws[f"A{row_num}"].alignment = Alignment(horizontal="left", vertical="center")

        ws[f"F{row_num}"] = service["rate"]
        ws[f"G{row_num}"] = service["quantity"]
        ws[f"H{row_num}"] = "=F{0}*G{0}".format(row_num)
        for col in ["F", "G", "H"]:
            cell = ws[f"{col}{row_num}"]
            cell.font = Font(name="Calibri", size=11, bold=False, italic=False)
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0" if col in {"F", "H"} else "0"
        _apply_full_table_border(ws, row_num, 1, 8)

    gross_row = 15 + len(active_rows)
    discount_row = gross_row + 1
    net_row = discount_row + 1

    if active_rows:
        gross_formula = f"=SUM(H15:H{gross_row - 1})"
    else:
        gross_formula = 0

    ws.merge_cells(f"A{gross_row}:E{gross_row}")
    ws[f"A{gross_row}"] = "Gross Consultation Charges"
    for col in ["F", "G"]:
        cell = ws[f"{col}{gross_row}"]
        cell.value = None
        cell.number_format = "#,##0"
        cell.font = Font(name="Calibri", size=11, bold=False, italic=False)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    _apply_summary_currency_style(ws, gross_row, gross_formula if active_rows else 0, percentage=False, bold=False)
    _apply_full_table_border(ws, gross_row, 1, 8)

    discount_rate = float(discount_percent) / 100.0
    ws.merge_cells(f"A{discount_row}:E{discount_row}")
    ws[f"A{discount_row}"] = "Friends and Family Discount %"
    for col in ["F", "G"]:
        cell = ws[f"{col}{discount_row}"]
        cell.value = None
        cell.number_format = "#,##0"
        cell.font = Font(name="Calibri", size=11, bold=False, italic=False)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    _apply_summary_currency_style(ws, discount_row, discount_rate, percentage=True, bold=False)
    _apply_full_table_border(ws, discount_row, 1, 8)

    ws.merge_cells(f"A{net_row}:E{net_row}")
    ws[f"A{net_row}"] = "Net Consultation Charges"
    ws[f"A{net_row}"].font = Font(name="Calibri", size=11, bold=True, italic=False)
    ws[f"A{net_row}"].alignment = Alignment(horizontal="left", vertical="center")
    for col in ["F", "G"]:
        cell = ws[f"{col}{net_row}"]
        cell.value = None
        cell.number_format = "#,##0"
        cell.font = Font(name="Calibri", size=11, bold=False, italic=False)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    ws[f"H{net_row}"] = f"=H{gross_row}*(1-H{discount_row})"
    ws[f"H{net_row}"].number_format = "#,##0"
    ws[f"H{net_row}"].font = Font(name="Calibri", size=11, bold=True, italic=False)
    ws[f"H{net_row}"].alignment = Alignment(horizontal="right", vertical="center")
    ws[f"H{net_row}"].border = Border(left=Side(style="thin", color="000000"), right=Side(style="thin", color="000000"), top=Side(style="thin", color="000000"), bottom=Side(style="thin", color="000000"))
    _apply_full_table_border(ws, net_row, 1, 8)

    signoff_row = net_row + 1
    signoff_text = "Thank you for your continued trust and support!"
    for col in range(1, 9):
        cell = ws.cell(row=signoff_row, column=col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = None if col != 1 else signoff_text
        cell.border = Border()
        cell.alignment = Alignment(horizontal="left", vertical="center")
        if col == 1:
            cell.font = Font(name="Calibri", size=11, bold=False, italic=True, color="000000")
            cell.alignment = Alignment(horizontal="left", vertical="center")
    return ws


def clear_unused_service_rows(ws, row_numbers=(15, 16, 17)):
    """Remove ghost service rows by clearing every cell and border in unused service slots."""
    _unmerge_row_ranges(ws, min(row_numbers), max(row_numbers))
    _safe_clear_cells(ws, min(row_numbers), max(row_numbers), 1, 8)
    return ws


def get_next_invoice_number(current_invoice: str) -> str:
    """
    Increments invoice numbers like '2627NL16' -> '2627NL17'.
    Rolls over '2627NL99' -> '2627NM01'.
    """
    match = re.match(r"^(\d{4})([A-Z]{2})(\d{2})$", current_invoice.strip())
    if not match:
        raise ValueError(f"Unexpected invoice format: {current_invoice}")

    year_prefix, letters, num_str = match.groups()
    num = int(num_str)

    if num < 99:
        return f"{year_prefix}{letters}{num + 1:02d}"

    c1, c2 = letters[0], letters[1]
    if c2 == "Z":
        if c1 == "N":
            next_letters = "PA"
        else:
            next_letters = chr(ord(c1) + 1) + "A"
    else:
        next_letters = c1 + chr(ord(c2) + 1)
    return f"{year_prefix}{next_letters}01"


def build_invoice_number(
    patient_name: str,
    month_label: str,
    patient_index: int = 1,
    prefix: str = SERIES_PREFIX,
    series_code: str = SERIES_CODE,
    start_sequence: int = SERIES_START,
) -> str:
    """Generate a strict chronological invoice number for the current FY series."""
    del patient_name, month_label
    if patient_index < 1:
        raise ValueError("patient_index must be >= 1")

    sequence = start_sequence + (patient_index - 1)
    if sequence <= 99:
        return f"{prefix}{series_code}{sequence:02d}"

    letters = list(series_code)
    while sequence > 99:
        sequence -= 99
        if letters[1] == "Z":
            if letters[0] == "N":
                letters = ["P", "A"]
            else:
                letters[0] = chr(ord(letters[0]) + 1)
                letters[1] = "A"
        else:
            letters[1] = chr(ord(letters[1]) + 1)

    return f"{prefix}{''.join(letters)}{sequence:02d}"
