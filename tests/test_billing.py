from datetime import date
from pathlib import Path

import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from app import (
    apply_exact_invoice_template,
    build_invoice_bundle_zip,
    build_monthly_summary_workbook,
    enrich_service_rows_with_discount,
    parse_discount_entries,
    parse_monthly_log,
    set_invoice_header,
    summarize_patient_totals,
)
from invoice_service import get_next_invoice_number
from tracker_service import ensure_monthly_tax_review_tracker


EXAMPLE_LOG = """Ravi - 7 (30)(PT)
Abhinay - 1 (60) (PT)+ 7(30)(TPR)
Punit - 7 (30)(TPR)
Shilpa - 7 (30)(TPR)
Gautham - 2(60)(TPR)
Rashmi - 2(45)(PT)
Kush-2(45)(TPR)
George -4(30)(TPR)
Simran - 6 (30) (TPR)
Shanta - 2 (30)(PT)
Dr kamala -7(30)(TPR)
Bhanu- 5 (30)(TPR)
Col Sushil - 1(30)(PT)
Anu-1 (30)(PT)
Naina- 4(30)(TPR)
Shivaram-1(30)(TPR)
Dr.arun kumar- 2(30)(TPR)
"""


def test_parse_monthly_log_and_totals():
    parsed_df, warnings = parse_monthly_log(EXAMPLE_LOG)

    assert not warnings
    assert len(parsed_df) > 0
    assert "Ravi" in set(parsed_df["patient"])

    patient_totals = summarize_patient_totals(parsed_df)
    abhinay = patient_totals[patient_totals["patient"] == "Abhinay"].iloc[0]

    assert abhinay["total_amount"] == 24750
    assert abhinay["total_sessions"] == 8

    rashmi = patient_totals[patient_totals["patient"] == "Rashmi"].iloc[0]
    assert rashmi["total_amount"] == 5000
    assert rashmi["total_sessions"] == 2

    assert patient_totals["total_amount"].sum() > 0


def test_parse_45_minute_rule():
    parsed_df, _ = parse_monthly_log("Rashmi - 2(45)(PT)\n")
    assert parsed_df.iloc[0]["duration_minutes"] == 30
    assert parsed_df.iloc[0]["amount"] == 5000


def test_apply_exact_invoice_template_for_multi_service_patient():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["I3"] = "OLD"
    ws["I4"] = "OLD"
    ws["C10"] = "OLD"
    ws["C11"] = "OLD"
    ws["C12"] = "OLD"
    for row in range(15, 18):
        ws[f"A{row}"] = "placeholder"
        ws[f"F{row}"] = 1
        ws[f"G{row}"] = 1
        ws[f"H{row}"] = 1
    ws["H18"] = "=SUM(H15:H17)"
    ws["H19"] = 0
    ws["H20"] = "=H18*(1-H19)"

    patient_df = pd.DataFrame(
        [
            {"patient": "Abhinay", "service_code": "PT", "session_count": 1, "duration_minutes": 60, "rate": 3750, "amount": 3750},
            {"patient": "Abhinay", "service_code": "PM", "session_count": 7, "duration_minutes": 30, "rate": 3000, "amount": 21000},
        ]
    )

    apply_exact_invoice_template(wb, "Abhinay", "31-Oct-2026", "2627NL16", patient_df)

    assert ws["I3"].value == "2627NL16"
    assert ws["I4"].value == date(2026, 10, 31)
    assert ws["I4"].number_format == "DD-MMM-YYYY"
    assert ws["C10"].value == "Abhinay"
    assert ws["C11"].value is None
    assert ws["C12"].value is None
    assert ws["A14"].value == "Description"
    assert ws["A15"].value == "Physical Medicine (30 min)"
    assert ws["A16"].value == "Rehabilitation (60 min)"
    assert ws["A17"].value == "Gross Consultation Charges"
    assert ws["F15"].value == 3000
    assert ws["G15"].value == 7
    assert ws["H15"].value == "=F15*G15"
    assert ws["F16"].value == 3750
    assert ws["G16"].value == 1
    assert ws["H16"].value == "=F16*G16"
    assert ws["H17"].value == "=SUM(H15:H16)"
    assert ws["H17"].number_format == "#,##0"
    assert ws["H17"].font.bold is False
    assert ws["A18"].value == "Friends and Family Discount %"
    assert ws["H18"].value == 0
    assert ws["H18"].number_format == "0%"
    assert ws["H18"].font.bold is False
    assert ws["A19"].value == "Net Consultation Charges"
    assert ws["H19"].value == "=H17*(1-H18)"
    assert ws["H19"].number_format == "#,##0"
    assert ws["H19"].font.bold is True
    assert ws["A6"].value is None
    assert ws["A11"].value is None
    assert ws["A12"].value is None
    assert ws.page_setup.orientation == ws.ORIENTATION_PORTRAIT
    assert str(ws.page_setup.paperSize) == ws.PAPERSIZE_A4
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 1


def test_header_removes_old_contact_rows_from_template():
    wb = Workbook()
    ws = wb.active
    ws.merge_cells("A10:B10")
    ws.merge_cells("C10:E10")
    ws.merge_cells("A11:B11")
    ws.merge_cells("C11:E11")
    ws.merge_cells("A12:B12")
    ws.merge_cells("C12:E12")
    ws["A10"] = "Client Name"
    ws["A11"] = "Phone number"
    ws["A12"] = "Email ID"
    ws["C10"] = "Shilpa"
    ws["C11"] = "123"
    ws["C12"] = "shilpa@example.com"

    set_invoice_header(ws, "Shilpa", "2627NL17", "2026-10-31")

    assert ws["A10"].value == "Client Name"
    assert ws["A11"].value is None
    assert ws["A12"].value is None
    assert ws["C10"].value == "Shilpa"
    assert ws["C11"].value is None
    assert ws["C12"].value is None
    assert "A10:B10" in {str(rng) for rng in ws.merged_cells.ranges}
    assert "C10:E10" in {str(rng) for rng in ws.merged_cells.ranges}
    assert "A11:B11" not in {str(rng) for rng in ws.merged_cells.ranges}
    assert "C11:E11" not in {str(rng) for rng in ws.merged_cells.ranges}


def test_parse_discount_entries_and_discounted_totals():
    discounts, warnings = parse_discount_entries("Abhinay - 20% discount\nShilpa - 25%\nInvalid line\n")
    assert len(warnings) == 1
    assert warnings[0].startswith("Line 3")
    assert discounts["abhinay"] == 20
    assert discounts["shilpa"] == 25

    parsed_df, _ = parse_monthly_log("Abhinay - 1(60)(PT)\nShilpa - 2(30)(PT)\n")
    patient_totals = summarize_patient_totals(parsed_df)
    discounted = patient_totals.copy()
    discounted["discount_percent"] = discounted["patient"].map({"Abhinay": 20, "Shilpa": 25}).fillna(0)
    discounted["net_amount"] = (discounted["total_amount"] * (1 - discounted["discount_percent"] / 100)).round().astype(int)

    abhinay_net = discounted.loc[discounted["patient"] == "Abhinay", "net_amount"].iloc[0]
    shilpa_net = discounted.loc[discounted["patient"] == "Shilpa", "net_amount"].iloc[0]
    assert abhinay_net == 3000
    assert shilpa_net == 3750


def test_service_detail_rows_include_gross_and_net_amounts():
    parsed_df, _ = parse_monthly_log("Shilpa - 7 (30)(TPR)\n")
    detail_df = enrich_service_rows_with_discount(parsed_df, {"shilpa": 25})

    assert "gross_amount" in detail_df.columns
    assert "net_amount" in detail_df.columns
    assert detail_df.iloc[0]["gross_amount"] == 21000
    assert detail_df.iloc[0]["net_amount"] == 15750


def test_single_service_patient_has_no_ghost_rows():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"

    patient_df = pd.DataFrame([
        {"patient": "Punit", "service_code": "PM", "session_count": 7, "duration_minutes": 30, "rate": 3000, "amount": 21000},
    ])

    apply_exact_invoice_template(wb, "Punit", "31-Oct-2026", "2627NL16", patient_df)

    assert ws["A14"].value == "Description"
    assert ws["A15"].value == "Physical Medicine (30 min)"
    assert ws["A16"].value == "Gross Consultation Charges"
    assert ws["A17"].value == "Friends and Family Discount %"
    assert ws["A18"].value == "Net Consultation Charges"
    assert ws["A19"].value == "Thank you for your continued trust and support!"
    assert ws["A19"].font.italic is True
    assert ws["A19"].font.name == "Calibri"
    assert ws["A19"].alignment.horizontal == "left"
    assert ws["A19"].border.left is None
    assert ws["A19"].border.right is None
    assert ws["A19"].border.top is None
    assert ws["A19"].border.bottom is None
    assert ws["H15"].value == "=F15*G15"
    assert ws["H16"].value == "=SUM(H15:H15)"
    assert ws["H18"].value == "=H16*(1-H17)"


def test_invoice_number_rollover_and_tracker_summary():
    assert get_next_invoice_number("2627NL16") == "2627NL17"
    assert get_next_invoice_number("2627NL99") == "2627NM01"
    assert get_next_invoice_number("2627NM99") == "2627NN01"
    assert get_next_invoice_number("2627NZ99") == "2627PA01"

    wb = ensure_monthly_tax_review_tracker(Path("data/test_monthly_review.xlsx"), "Oct 2026", [
        {"patient": "Abhinay", "invoice_number": "2627NL16", "amount": 24750, "total_sessions": 8, "service_summary": "PT, PM", "review_status": "Reviewed"},
        {"patient": "Ravi", "invoice_number": "2627NL17", "amount": 17500, "total_sessions": 7, "service_summary": "PT", "review_status": "Approved"},
    ])
    ws = wb["Monthly Tax & Review"]
    assert ws["A2"].value == "Oct 2026"
    assert ws["C2"].value == 42250
    assert ws["D2"].value == 42250
    assert ws["E2"].value == 15
    assert ws["F2"].value == "PT, PM; PT"


def test_monthly_summary_and_bundle_export(tmp_path):
    summary_path = tmp_path / "Monthly_Summary_Oct_2026.xlsx"
    bundle_path = tmp_path / "Invoice_Bundle_Oct_2026.zip"
    invoice_rows = [
        {"patient": "Abhinay", "invoice_number": "2627NL16", "amount": 24750, "total_sessions": 8, "service_summary": "PT, PM", "review_status": "Reviewed", "invoice_date": "2026-10-31"},
        {"patient": "Ravi", "invoice_number": "2627NL17", "amount": 17500, "total_sessions": 7, "service_summary": "PT", "review_status": "Approved", "invoice_date": "2026-10-31"},
    ]

    build_monthly_summary_workbook("Oct 2026", invoice_rows, summary_path)
    bundle_file = tmp_path / "invoice_a.xlsx"
    bundle_file.write_bytes(b"excel")
    pdf_file = tmp_path / "invoice_a.pdf"
    pdf_file.write_bytes(b"pdf")
    build_invoice_bundle_zip([bundle_file, pdf_file], summary_path, bundle_path)

    assert summary_path.exists()
    assert bundle_path.exists()
    with zipfile.ZipFile(bundle_path) as archive:
        names = set(archive.namelist())
        assert "invoice_a.xlsx" in names
        assert "invoice_a.pdf" in names
        assert "Monthly_Summary_Oct_2026.xlsx" in names
