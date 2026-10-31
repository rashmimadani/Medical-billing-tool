from pathlib import Path

from openpyxl import Workbook, load_workbook


def ensure_monthly_tax_review_tracker(tracker_path, month_label: str, invoice_rows):
    tracker_path = Path(tracker_path)
    if tracker_path.exists():
        wb = load_workbook(tracker_path)
        if "Monthly Tax & Review" not in wb.sheetnames:
            ws = wb.create_sheet("Monthly Tax & Review")
        else:
            ws = wb["Monthly Tax & Review"]
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Monthly Tax & Review"

    headers = [
        "Month",
        "Patients",
        "Gross Amount",
        "Net Amount",
        "Sessions",
        "Service Summary",
        "Reviewed",
        "Status",
    ]
    existing_headers = [ws.cell(row=1, column=idx).value for idx in range(1, ws.max_column + 1)] if ws.max_row > 0 else []
    if existing_headers and ("Tax Amount" in existing_headers or "Approved" in existing_headers or existing_headers[: len(headers)] != headers):
        ws.delete_rows(1, ws.max_row)
        existing_headers = []

    if not existing_headers:
        for idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=idx, value=header)
        summary_row = 2
    else:
        summary_row = 2

    gross_total = sum(int(row.get("gross_amount", row.get("amount", 0))) for row in invoice_rows)
    net_total = sum(int(row.get("net_amount", row.get("amount", 0))) for row in invoice_rows)
    total_sessions = sum(int(row.get("total_sessions", 0)) for row in invoice_rows)
    reviewed_count = sum(1 for row in invoice_rows if str(row.get("review_status", "")).lower() in {"reviewed", "approved"})
    service_summary = "; ".join(
        str(row.get("service_summary", "")).strip()
        for row in invoice_rows
        if str(row.get("service_summary", "")).strip()
    )

    ws[f"A{summary_row}"] = month_label
    ws[f"B{summary_row}"] = len(invoice_rows)
    ws[f"C{summary_row}"] = gross_total
    ws[f"D{summary_row}"] = net_total
    ws[f"E{summary_row}"] = total_sessions
    ws[f"F{summary_row}"] = service_summary
    ws[f"G{summary_row}"] = reviewed_count
    ws[f"H{summary_row}"] = "Ready for review" if invoice_rows else "No data"

    for row_number, row in enumerate(invoice_rows, start=3):
        ws[f"A{row_number}"] = month_label
        ws[f"B{row_number}"] = row.get("patient", "")
        ws[f"C{row_number}"] = int(row.get("gross_amount", row.get("amount", 0)))
        ws[f"D{row_number}"] = int(row.get("net_amount", row.get("amount", 0)))
        ws[f"E{row_number}"] = int(row.get("total_sessions", 0))
        ws[f"F{row_number}"] = row.get("service_summary", "")
        ws[f"G{row_number}"] = row.get("review_status", "Reviewed")
        ws[f"H{row_number}"] = "Ready for review"

    tracker_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(tracker_path)
    return wb
